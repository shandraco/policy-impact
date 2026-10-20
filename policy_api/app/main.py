"""Policy API. The only path from Hermes to data, models and decisions.

Internal routes (never exposed by the proxy):
  POST /gate/v1/chat/completions   the model gate; Hermes's model provider points here
  POST /tools/{name}               the five tools; called by the Hermes plugin
  POST /check-reply                Open WebUI's outlet filter checks the model's reply
Routes the proxy exposes to browsers:
  /review/...                      the review page
  GET /exports/{run}.xlsx          workbook download, by signed link
"""
from __future__ import annotations

import hmac
import io
import logging
from contextlib import asynccontextmanager
from urllib.parse import parse_qs

from fastapi import Depends, FastAPI, File, Header, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from starlette.middleware.trustedhost import TrustedHostMiddleware

from . import review
from .answer import md
from .config import load_settings
from .db import all_rows, audit, one_row
from .gate import Gate, as_sse
from .ingest import store_version
from .links import session_token
from .pipeline import Services, build_services
from .safe_parse import ParseTimeout, parse_isolated
from .scope import Scope, ScopeError, verify
from .tools import check_reply, dispatch

log = logging.getLogger(__name__)
SETTINGS = load_settings()
COOKIE = "pi_review"
MAX_UPLOAD = 20 * 1024 * 1024


@asynccontextmanager
async def lifespan(app: FastAPI):
    for name in ("scope_signing_key", "tool_key", "gate_key"):
        if len(getattr(SETTINGS, name)) < 16:
            raise RuntimeError(f"{name.upper()} must be set to a long random value")
    svc = await build_services(SETTINGS)
    app.state.svc = svc
    app.state.gate = Gate(svc)
    yield
    await app.state.gate.aclose()
    await svc.close()


app = FastAPI(title="Policy API", lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=list(SETTINGS.allowed_hosts))


def svc_of(request: Request) -> Services:
    return request.app.state.svc


def _bearer(expected: str):
    def check(authorization: str = Header(default="")) -> None:
        if not hmac.compare_digest(authorization.removeprefix("Bearer ").strip(), expected):
            raise HTTPException(401, "Wrong key")
    return check


@app.get("/health")
async def health(request: Request) -> dict:
    async with svc_of(request).db.system() as conn:
        stats = await review.queue_stats(conn)
        jobs = await one_row(conn, "SELECT count(*) FILTER (WHERE status = 'dead') AS dead, count(*) FILTER (WHERE status = 'pending') AS waiting FROM jobs")
    return {"ok": True, "config": svc_of(request).config_hash, "review_queue": stats, "jobs": jobs}


# --- model gate -----------------------------------------------------------------------------------------

@app.get("/gate/v1/models", dependencies=[Depends(_bearer(SETTINGS.gate_key))])
async def gate_models() -> dict:
    return {"object": "list", "data": [{"id": SETTINGS.chat_model, "object": "model", "owned_by": "policy-api"}]}


@app.post("/gate/v1/chat/completions", dependencies=[Depends(_bearer(SETTINGS.gate_key))])
async def gate_chat(request: Request):
    body = await request.json()
    resp = await request.app.state.gate.chat(body)
    if body.get("stream"):
        return StreamingResponse(iter(as_sse(resp)), media_type="text/event-stream")
    return resp


# --- tools ----------------------------------------------------------------------------------------------

@app.post("/tools/{name}", dependencies=[Depends(_bearer(SETTINGS.tool_key))])
async def tool(request: Request, name: str) -> dict:
    args = await request.json()
    if not isinstance(args, dict):
        raise HTTPException(400, "Arguments must be an object")
    try:
        scope = verify(SETTINGS.scope_signing_key, str(args.pop("scope_token", "")))
    except ScopeError as e:
        return {"display": f"I can't confirm who you are ({e}). Reload the page and try again.", "facts": {"ids": []}}
    return dict(await dispatch(svc_of(request), scope, name, args))


@app.post("/check-reply")
async def check(request: Request) -> dict:
    body = await request.json()
    try:
        scope = verify(SETTINGS.scope_signing_key, str(body.get("token", "")))
    except ScopeError:
        raise HTTPException(401, "Bad token") from None
    return await check_reply(svc_of(request), scope, str(body.get("reply", "")))


# --- review page ----------------------------------------------------------------------------------------

def _scope_from_cookie(request: Request) -> Scope:
    try:
        return verify(SETTINGS.scope_signing_key, request.cookies.get(COOKIE, ""))
    except ScopeError:
        raise HTTPException(401, "Open the review page from the link in the chat.") from None


def _html(text: str, status: int = 200) -> HTMLResponse:
    return HTMLResponse(text, status, headers={"Cache-Control": "no-store", "X-Frame-Options": "DENY",
                                               "Content-Security-Policy": "default-src 'none'; style-src 'unsafe-inline'; form-action 'self'"})


@app.get("/review/login")
async def review_login(t: str) -> RedirectResponse:
    try:
        scope = verify(SETTINGS.scope_signing_key, t)
    except ScopeError:
        return _html(review.page("Link expired", "<h1>This link has expired</h1><p>Ask for the review page again in the chat.</p>"), 401)
    r = RedirectResponse("/review", 303)
    r.set_cookie(COOKIE, session_token(SETTINGS, scope), httponly=True, samesite="strict",
                 secure=SETTINGS.public_base_url.startswith("https"), max_age=8 * 3600, path="/review")
    return r


@app.get("/review")
async def review_home(request: Request):
    scope = _scope_from_cookie(request)
    if not scope.can("review") and not scope.can("upload"):
        return _html(review.page("Review", "<h1>No access</h1><p>Your role can't use the review page.</p>"), 403)
    return _html(await review.home(svc_of(request), scope))


@app.get("/review/run/{run_id}")
async def review_run(request: Request, run_id: str):
    scope = _scope_from_cookie(request)
    scope.can("review") or _forbid()
    page = await review.run_form(svc_of(request), scope, run_id)
    return _html(page) if page else RedirectResponse("/review", 303)


@app.post("/review/run/{run_id}")
async def review_run_confirm(request: Request, run_id: str):
    scope = _scope_from_cookie(request)
    form = parse_qs((await request.body()).decode())
    await review.confirm_run(svc_of(request), scope, run_id, form)
    return _html(review.page("Confirmed", "<h1>Confirmed</h1><p>Every entity is being checked now. Ask in the chat for the results in a few minutes.</p>"))


@app.get("/review/next")
async def review_next(request: Request, saved: str | None = None, after: int = 0):
    scope = _scope_from_cookie(request)
    scope.can("review") or _forbid()
    return _html(await review.next_finding(svc_of(request), scope, saved, after))


@app.post("/review/finding/{finding_id}")
async def review_decide(request: Request, finding_id: int):
    scope = _scope_from_cookie(request)
    decision = parse_qs((await request.body()).decode()).get("decision", [""])[0]
    ok = await review.decide(svc_of(request), scope, finding_id, decision)
    return RedirectResponse(f"/review/next?saved={finding_id}" if ok else "/review/next", 303)


@app.post("/review/undo/{finding_id}")
async def review_undo(request: Request, finding_id: int):
    scope = _scope_from_cookie(request)
    await review.undo(svc_of(request), scope, finding_id)
    return RedirectResponse(f"/review/next?after={finding_id - 1}", 303)


@app.post("/review/upload")
async def review_upload(request: Request, file: UploadFile = File(...)):
    scope = _scope_from_cookie(request)
    scope.require("upload")
    svc = svc_of(request)
    data = await file.read(MAX_UPLOAD + 1)
    back = "<p><a href='/review'>Back to review home</a></p>"
    if len(data) > MAX_UPLOAD:
        return _html(review.page("Too big", f"<h1>That file is too big</h1><p>The limit is 20 MB.</p>{back}"), 413)
    try:
        doc = await parse_isolated(file.filename or "upload", data, SETTINGS.policy_id_pattern)
    except (ValueError, ParseTimeout) as e:
        return _html(review.page("Can't read", f"<h1>That file couldn't be read</h1><p>{md(e) or 'It took too long.'}</p>{back}"), 400)
    async with svc.db.as_person(scope) as conn:
        pol = await one_row(conn, "SELECT id, entity_id FROM policies WHERE id = %s", doc.id or "")
    if not pol:
        return _html(review.page("Not found", f"<h1>No matching policy</h1><p>The file's ID ({md(doc.id or 'none')}) doesn't match a "
                                 f"policy you can see. Put the ID at the top, for example <code>id: DOE-512</code>.</p>{back}"), 400)
    import hashlib
    version = await store_version(svc.db, svc.embedder, doc, pol["entity_id"], hashlib.sha256(data).hexdigest(),
                                  SETTINGS.policy_id_pattern, source="draft")
    async with svc.db.system() as conn:
        await audit(conn, scope.user_id, scope.role, "upload_draft", {"policy": doc.id, "version": version})
    return _html(review.page("Uploaded", f"<h1>Draft uploaded</h1><p>{md(doc.id)} version {version} is saved as a draft. "
                             f"Ask in the chat: <b>how does the {md(doc.id)} change affect the entities below?</b></p>{back}"))


def _forbid():
    raise HTTPException(403, "Your role can't do this.")


# --- downloads ------------------------------------------------------------------------------------------

@app.get("/exports/{run_id}.xlsx")
async def export(request: Request, run_id: str, t: str):
    try:
        scope = verify(SETTINGS.scope_signing_key, t)
        scope.require("export")
    except ScopeError:
        raise HTTPException(403, "This link has expired. Ask for the results again in the chat.") from None
    from openpyxl import Workbook
    svc = svc_of(request)
    async with svc.db.as_person(scope) as conn:  # the workbook holds only rows this person may see
        run = await one_row(conn, "SELECT * FROM runs WHERE id = %s", run_id)
        if not run:
            raise HTTPException(404, "Run not found")
        rows = await all_rows(conn, """SELECT e.name AS entity, f.requirement, f.result, c.policy_id, c.number, f.why, f.evidence, f.stage
                                       FROM findings f JOIN entities e ON e.id = f.entity_id LEFT JOIN clauses c ON c.id = f.clause_id
                                       WHERE f.run_id = %s ORDER BY e.name, f.requirement""", run_id)
        await audit(conn, scope.user_id, scope.role, "download", {"run_id": run_id, "rows": len(rows)})
    reqs = run["requirements"] or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    ws.append(["Entity", "Result", "Their policy", "Section", "Requirement", "Why", "Evidence", "Decided by"])
    for r in rows:
        req = reqs[r["requirement"]]["text"] if r["requirement"] < len(reqs) else ""
        ws.append([r["entity"], r["result"].replace("_", " "), r["policy_id"] or "", r["number"] or "", req, r["why"], r["evidence"],
                   "a reviewer" if r["stage"] == "reviewer" else ("code rule" if r["stage"] == "rule" else "model")])
    rq = wb.create_sheet("Requirements")
    rq.append(["#", "Section", "Requirement", "Checked"])
    for i, r in enumerate(reqs, 1):
        rq.append([i, r["source"], r["text"], "yes" if r.get("keep") else "no"])
    for sheet in (ws, rq):
        for row in sheet.iter_rows():  # policy text starting with "=" must stay text, never become a formula
            for c in row:
                if c.data_type == "f":
                    c.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return StreamingResponse(iter([buf.getvalue()]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{run["policy_id"]}-impact.xlsx"'})


@app.exception_handler(ScopeError)
async def scope_error(request: Request, exc: ScopeError) -> JSONResponse:
    return JSONResponse({"detail": str(exc)}, 403)
